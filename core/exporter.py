"""
数据导出模块 - 提供多种格式的数据导出功能
"""

import io
import csv
import logging
from typing import Optional, List, Any, Dict
from enum import Enum
import numpy as np

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    logging.getLogger(__name__).warning("openpyxl 未安装，Excel 导出功能不可用")

logger = logging.getLogger(__name__)


class ExportFormat(str, Enum):
    """导出格式枚举"""
    CSV = "csv"
    XLSX = "xlsx"
    TXT = "txt"
    NPY = "npy"
    MAT = "mat"


class DataExporter:
    """
    数据导出器
    
    支持格式:
    - CSV: 逗号分隔值
    - XLSX: Excel 格式
    - TXT: 文本格式
    - NPY: NumPy 格式
    - MAT: MATLAB 格式
    
    Methods:
        export: 导出数据
        export_to_csv: 导出为 CSV
        export_to_excel: 导出为 Excel
        export_to_text: 导出为文本
        export_to_npy: 导出为 NPY
        export_to_mat: 导出为 MAT
    """
    
    def __init__(self):
        """初始化导出器"""
        logger.debug("DataExporter 初始化")
    
    def export(self, data: Any, format: str, **kwargs) -> bytes:
        """
        导出数据
        
        Args:
            data: 要导出的数据
            format: 导出格式
            **kwargs: 额外参数
            
        Returns:
            导出的二进制数据
            
        Raises:
            ValueError: 不支持的格式
        """
        format = format.lower()
        
        if format == ExportFormat.CSV:
            return self.export_to_csv(data, **kwargs)
        elif format == ExportFormat.XLSX:
            return self.export_to_excel(data, **kwargs)
        elif format == ExportFormat.TXT:
            return self.export_to_text(data, **kwargs)
        elif format == ExportFormat.NPY:
            return self.export_to_npy(data, **kwargs)
        elif format == ExportFormat.MAT:
            return self.export_to_mat(data, **kwargs)
        else:
            raise ValueError(f"不支持的导出格式：{format}")
    
    def export_to_csv(self, data: Any, header: Optional[str] = None,
                     row_index: bool = False, trim_nulls: bool = False,
                     row_start: int = 0, row_end: Optional[int] = None,
                     col_start: int = 0, col_end: Optional[int] = None) -> bytes:
        """
        导出为 CSV 格式
        
        Args:
            data: 数据（二维列表或数组）
            header: 表头模板（可选）
            row_index: 是否添加行号
            trim_nulls: 是否裁剪 null/nan/空值
            row_start: 起始行
            row_end: 结束行
            col_start: 起始列
            col_end: 结束列
            
        Returns:
            CSV 二进制数据
        """
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        
        # 处理数据范围
        if isinstance(data, (list, np.ndarray)):
            data = self._slice_data(data, row_start, row_end, col_start, col_end)
        
        # 生成表头
        if header:
            headers = self._generate_headers(header, data, row_index)
            writer.writerow(headers)
        
        # 写入数据
        if isinstance(data, np.ndarray):
            data = data.tolist()
        
        for i, row in enumerate(data):
            if row_index:
                row = [i] + list(row)
            
            if trim_nulls:
                row = self._trim_row(row)
            
            writer.writerow(row)
        
        return buffer.getvalue().encode('utf-8-sig')
    
    def export_to_excel(self, data: Any, header: Optional[str] = None,
                       row_index: bool = False, trim_nulls: bool = False,
                       row_start: int = 0, row_end: Optional[int] = None,
                       col_start: int = 0, col_end: Optional[int] = None) -> bytes:
        """
        导出为 Excel 格式
        
        Args:
            data: 数据
            header: 表头模板
            row_index: 是否添加行号
            trim_nulls: 是否裁剪 null/nan/空值
            row_start: 起始行
            row_end: 结束行
            col_start: 起始列
            col_end: 结束列
            
        Returns:
            Excel 二进制数据
        """
        if not HAS_OPENPYXL:
            raise ImportError("openpyxl 未安装")
        
        buffer = io.BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        
        # 处理数据范围
        if isinstance(data, (list, np.ndarray)):
            data = self._slice_data(data, row_start, row_end, col_start, col_end)
        
        # 生成表头
        row_offset = 0
        if header:
            headers = self._generate_headers(header, data, row_index)
            for col_idx, h in enumerate(headers, 1):
                ws.cell(row=1, column=col_idx, value=h)
            row_offset = 1
        
        # 写入数据
        if isinstance(data, np.ndarray):
            data = data.tolist()
        
        for i, row in enumerate(data):
            if row_index:
                row = [i] + list(row)
            
            if trim_nulls:
                row = self._trim_row(row)
            
            for col_idx, val in enumerate(row, 1):
                ws.cell(row=row_offset + i + 1, column=col_idx, value=val)
        
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    
    def export_to_text(self, data: Any, delimiter: str = "\t",
                      row_start: int = 0, row_end: Optional[int] = None,
                      col_start: int = 0, col_end: Optional[int] = None) -> bytes:
        """
        导出为文本格式
        
        Args:
            data: 数据
            delimiter: 分隔符
            row_start: 起始行
            row_end: 结束行
            col_start: 起始列
            col_end: 结束列
            
        Returns:
            文本二进制数据
        """
        buffer = io.StringIO()
        
        # 处理数据范围
        if isinstance(data, (list, np.ndarray)):
            data = self._slice_data(data, row_start, row_end, col_start, col_end)
        
        if isinstance(data, np.ndarray):
            data = data.tolist()
        
        for row in data:
            buffer.write(delimiter.join(str(v) for v in row))
            buffer.write("\n")
        
        return buffer.getvalue().encode('utf-8-sig')
    
    def export_to_npy(self, data: Any) -> bytes:
        """
        导出为 NumPy NPY 格式
        
        Args:
            data: NumPy 数组或可转换为数组的数据
            
        Returns:
            NPY 二进制数据
        """
        buffer = io.BytesIO()
        
        if not isinstance(data, np.ndarray):
            data = np.array(data)
        
        np.save(buffer, data)
        buffer.seek(0)
        return buffer.getvalue()
    
    def export_to_mat(self, data: Any, var_name: str = "data") -> bytes:
        """
        导出为 MATLAB MAT 格式
        
        Args:
            data: 数据
            var_name: 变量名
            
        Returns:
            MAT 二进制数据
        """
        buffer = io.BytesIO()
        
        if not isinstance(data, np.ndarray):
            data = np.array(data)
        
        scipy.io.savemat(buffer, {var_name: data})
        buffer.seek(0)
        return buffer.getvalue()
    
    def _slice_data(self, data: Any, row_start: int, row_end: Optional[int],
                   col_start: int, col_end: Optional[int]) -> Any:
        """切片数据"""
        if isinstance(data, np.ndarray):
            return data[row_start:row_end, col_start:col_end]
        elif isinstance(data, list):
            return [row[col_start:col_end] for row in data[row_start:row_end]]
        return data
    
    def _generate_headers(self, template: str, data: Any, row_index: bool) -> List[str]:
        """
        根据模板生成表头
        
        Args:
            template: 表头模板
            data: 数据
            row_index: 是否包含行号列
            
        Returns:
            表头列表
        """
        if not data:
            return []
        
        num_cols = len(data[0]) if isinstance(data[0], (list, np.ndarray)) else 1
        
        headers = []
        
        if row_index:
            headers.append("Row_num")
        
        for i in range(num_cols):
            header = template.replace("%Col_num", str(i)).replace("%Row_num", str(i))
            header = header.replace("%VarName", "Data")
            headers.append(header)
        
        return headers
    
    def _trim_row(self, row: List[Any]) -> List[Any]:
        """裁剪行中的 null/nan/空值"""
        trimmed = []
        for val in row:
            if val is None:
                continue
            if isinstance(val, float) and (np.isnan(val) or np.isinf(val)):
                continue
            if isinstance(val, str) and val.strip() == "":
                continue
            trimmed.append(val)
        return trimmed if trimmed else row
